import copy
import io
import json
from unittest.mock import AsyncMock

import pytest
from openpyxl import Workbook, load_workbook

from app.utils.excel_dict import ExcelDictConfig, ExcelDictSnapshot
from app.utils.excel_util import ExcelUtil


def node(key, parent, code, label, status=1):
    return {"id": key, "parent_id": parent, "dict_code": code, "dict_name": label, "status": status}


FLAT = [node(1, None, "test", "类型"), node(2, 1, "01", "选项一"), node(3, 1, "02", "选项二")]
TREE = [node(1, None, "test", "类型"), node(2, 1, "east", "东区"), node(3, 1, "west", "西区"), node(4, 2, "01", "一组"), node(5, 3, "01", "一组")]


def snapshot(config, rows=None):
    return ExcelDictSnapshot({"value": config}, {"test": rows or FLAT})


@pytest.mark.parametrize("storage,value", [("array", ["01", "02", "01"]), ("json", '["01", "02", "01"]'), ("delimited", "01;02;01")])
@pytest.mark.parametrize("excel_format", ["text", "json"])
def test_multiple_formats_round_trip(storage, value, excel_format):
    converter = snapshot(ExcelDictConfig("test", mode="multiple", storage=storage, storage_separator=";", separator="|", excel_format=excel_format))
    original = copy.deepcopy(value)
    exported = converter.convert("value", value)
    assert converter.convert("value", exported, importing=True) == value
    for text in ('["选项一","选项二","选项一"]', "选项一|选项二|选项一"):
        assert converter.convert("value", text, importing=True) == value
    assert value == original


@pytest.mark.parametrize(
    "mode,value,text",
    [
        ("cascade", ["east", "01"], "东区 / 一组"),
        ("cascade_multiple", [["east", "01"], ["west", "01"]], "东区 / 一组,西区 / 一组"),
    ],
)
@pytest.mark.parametrize("storage", ["array", "json"])
def test_cascade_round_trip(mode, value, text, storage):
    converter = snapshot(ExcelDictConfig("test", mode=mode, storage=storage), TREE)
    expected = json.dumps(value, ensure_ascii=False) if storage == "json" else value
    assert converter.convert("value", expected) == text
    assert converter.convert("value", text, importing=True) == expected
    labels = ["东区", "一组"] if mode == "cascade" else [["东区", "一组"], ["西区", "一组"]]
    assert converter.convert("value", json.dumps(labels), importing=True) == expected


@pytest.mark.parametrize("value", ["一组", "东区 / 缺失", '["东区", ["一组"]]', "西区 / 东区"])
def test_invalid_paths_have_context(value):
    converter = snapshot(ExcelDictConfig("test", mode="cascade"), TREE)
    _, errors = converter.import_row({"value": value}, 7, {"value": "区域"})
    assert len(errors) == 1 and "第7行" in errors[0] and "区域" in errors[0] and repr(value) in errors[0]


@pytest.mark.parametrize("value", ['["east"]', '["east", ["01"]]', "东区 / 一组"])
def test_cascade_multiple_requires_nested_paths(value):
    converter = snapshot(ExcelDictConfig("test", mode="cascade_multiple"), TREE)
    if value == "东区 / 一组":
        assert converter.convert("value", value, importing=True) == [["east", "01"]]
    else:
        with pytest.raises(ValueError):
            converter.convert("value", value, importing=True)


@pytest.mark.parametrize("code,kind,expected", [("01", str, "01"), ("01", int, 1), ("true", bool, True), ("false", bool, False)])
def test_typed_codes(code, kind, expected):
    converter = snapshot(ExcelDictConfig("test", code_type=kind), [FLAT[0], node(2, 1, code, "选项")])
    assert converter.convert("value", "选项", importing=True) == expected
    assert converter.convert("value", code) == "选项"


def test_unknown_empty_duplicate_labels_and_delimiters():
    converter = snapshot(ExcelDictConfig("test"))
    for value in (None, ""):
        assert converter.convert("value", value, importing=True) == value
    assert converter.convert("value", "unknown") == "unknown"
    with pytest.raises(ValueError):
        converter.convert("value", "未知", importing=True)
    converter = snapshot(ExcelDictConfig("test"), [*FLAT, node(4, 1, "03", "选项一")])
    with pytest.raises(ValueError, match="歧义"):
        converter.convert("value", "选项一", importing=True)
    converter = snapshot(ExcelDictConfig("test", mode="multiple"), [FLAT[0], node(2, 1, "01", "选项,一")])
    text = converter.convert("value", ["01"])
    assert json.loads(text) == ["选项,一"]
    assert converter.convert("value", text, importing=True) == ["01"]
    assert json.loads(converter.convert("value", ["unknown"])) == ["unknown"]
    with pytest.raises(ValueError):
        converter.convert("value", '["选项,一"', importing=True)


async def test_snapshot_loads_each_type_once_and_filters_import():
    service = AsyncMock()
    service.all_nodes.return_value = copy.deepcopy(FLAT) + [node(4, 1, "03", "已停用", 0)]
    service.effective.return_value = FLAT
    configs = {"a": ExcelDictConfig("test"), "b": ExcelDictConfig("test")}
    exported = await ExcelDictSnapshot.load(service, configs)
    assert exported.convert("a", "03") == "已停用"
    service.all_nodes.assert_awaited_once_with("test")
    imported = await ExcelDictSnapshot.load(service, configs, importing=True)
    with pytest.raises(ValueError):
        imported.convert("a", "已停用", importing=True)
    service.effective.assert_awaited_once_with("test", None)
    service.all_nodes.return_value[1]["dict_name"] = "修改"
    assert exported.convert("a", "01") == "选项一"


def workbook_bytes(rows):
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


async def test_excel_integration_legacy_and_physical_row_numbers():
    converter = snapshot(ExcelDictConfig("test", mode="multiple"))
    data = [{"value": ["01", "02"]}]
    original = copy.deepcopy(data)
    output = await ExcelUtil.aexport_list2excel(data, {"value": "选项"}, dictionary=converter)
    assert data == original
    assert await ExcelUtil.aread_excel_to_dicts(output, dictionary=converter, header_mapping={"选项": "value"}) == data
    with pytest.raises(ValueError, match="第4行"):
        await ExcelUtil.aread_excel_to_dicts(workbook_bytes([["选项"], ["选项一"], [None], ["未知"]]), dictionary=converter, header_mapping={"选项": "value"})
    output = await ExcelUtil.aexport_list2excel([{"id": 1}], {"id": "编号"})
    assert ExcelUtil.read_excel_to_dicts(output) == [{"编号": 1}]


def test_template_options_and_instructions():
    configs = {"value": ExcelDictConfig("test"), "multi": ExcelDictConfig("test", mode="multiple", separator=";")}
    converter = ExcelDictSnapshot(configs, {"test": FLAT})
    output = ExcelUtil.get_excel_template(["选项", "多选"], [], [], dictionary=converter, mapping_dict={"value": "选项", "multi": "多选"})
    wb = load_workbook(io.BytesIO(output))
    assert wb["字典选项"].sheet_state == "hidden"
    assert wb["字典选项"]["A1"].value == "选项一"
    assert "选项一;选项二" in wb.active["B1"].comment.text
    assert len(wb.active.data_validations.dataValidation) == 1
    wb.close()


def test_custom_cascade_separators_and_empty_arrays():
    converter = snapshot(ExcelDictConfig("test", mode="cascade_multiple", separator=";", path_separator=" > "), TREE)
    value = [["east", "01"], ["west", "01"]]
    assert converter.convert("value", value) == "东区 > 一组;西区 > 一组"
    assert converter.convert("value", "东区 > 一组;西区 > 一组", importing=True) == value
    for mode in ("multiple", "cascade", "cascade_multiple"):
        converter = snapshot(ExcelDictConfig("test", mode=mode), TREE)
        assert converter.convert("value", []) == "[]"
        assert converter.convert("value", "[]", importing=True) == []


@pytest.mark.parametrize(
    "kwargs",
    [
        {"separator": ""},
        {"path_separator": ""},
        {"storage_separator": ""},
        {"mode": "cascade", "storage": "delimited"},
        {"code_type": float},
        {"mode": "cascade_multiple", "separator": "/", "path_separator": " / "},
    ],
)
def test_invalid_configuration(kwargs):
    with pytest.raises(ValueError):
        ExcelDictConfig("test", **kwargs)
