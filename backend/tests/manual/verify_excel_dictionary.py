"""Opt-in local integration check. Run from backend/ with RESNEXUS_TEST_PASSWORD set.

Uses the existing development docs login flow, Vite proxy, and local MySQL.
Creates only a uniquely named disabled user without roles, then soft-deletes it.
"""

import asyncio
import io
import logging
import os
import sys
import time
from pathlib import Path

import httpx
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

BASE = "http://127.0.0.1:5180/api/v1"
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def data(response):
    response.raise_for_status()
    body = response.json()
    assert body.get("success", True), body.get("msg")
    return body.get("data", body)


async def verify_database(username):
    from sqlalchemy import text

    from app.config.setting import settings
    from app.core.database import databases
    from app.utils.password_util import PwdUtil

    assert settings.DATABASE_HOST in ("127.0.0.1", "localhost"), "Local database only"
    try:
        async with databases.session() as session:
            result = await session.execute(text("SELECT gender, status, password, is_deleted FROM sys_user WHERE username = :username"), {"username": username})
            row = result.mappings().one()
            assert row["gender"] == "1" and row["status"] == 0 and not row["is_deleted"]
            assert await PwdUtil.averify_password(settings.PASSWORD_IMPORT_DEFAULT, row["password"])
    finally:
        await databases.dispose()


def filled_template(template, username, gender):
    book = load_workbook(io.BytesIO(template))
    book.active.append([1, username, "Excel字典联调测试", None, None, gender, "停用"])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def main():
    for logger in ("httpx", "httpcore"):
        logging.getLogger(logger).setLevel(logging.WARNING)
    password = os.environ["RESNEXUS_TEST_PASSWORD"]
    name = f"excel_live_{time.time_ns()}"[:32]
    with httpx.Client(base_url=BASE, timeout=45) as client:
        health = data(client.get("/monitor/health/check"))
        assert health["environment"] == "dev" and health["db_status"] == health["redis_status"] == 1
        login = data(client.post("/system/auth/login", data={"username": "super", "password": password}, headers={"Referer": BASE + "/docs"}))
        client.headers["Authorization"] = "Bearer " + login["access_token"]
        template = client.get("/system/user/import/template")
        template.raise_for_status()
        book = load_workbook(io.BytesIO(template.content))
        assert [cell.value for cell in book.active[1]] == ["部门编号", "账号", "昵称", "邮箱", "手机号", "性别", "状态"]
        assert {book["字典选项"].cell(i, 6).value for i in range(1, 4)} == {"男", "女", "未知"}
        assert {book["字典选项"].cell(i, 7).value for i in range(1, 3)} == {"启用", "停用"}
        assert len(book.active.data_validations.dataValidation) == 2
        book.close()
        print("PASS: Vite proxy, MySQL/Redis health, authenticated template and dictionary options")
        try:
            result = data(client.post("/system/user/import/data", files={"file": ("valid.xlsx", filled_template(template.content, name, "女"), MIME)}))
            assert result == "成功导入 1 条数据", result
            # Deliberately no retry/sleep: success must mean data is already committed.
            items = data(client.get("/system/user/list", params={"username": name}))["items"]
            assert len(items) == 1 and items[0]["username"] == name
            assert items[0]["gender"] == "1" and items[0]["status"] == 0
            assert not items[0]["is_superuser"] and items[0]["role_ids"] == []
            asyncio.run(verify_database(name))
            print("PASS: import, immediate list read, committed MySQL codes and password hash")

            exported = client.post("/system/user/export", json={"username": name})
            exported.raise_for_status()
            book = load_workbook(io.BytesIO(exported.content))
            rows = list(book.active.values)
            assert len(rows) == 2
            row = dict(zip(rows[0], rows[1], strict=True))
            assert row["性别"] == "女" and row["状态"] == "停用" and row["是否超级管理员"] == "否"
            book.close()
            print("PASS: filtered export restores dictionary labels")

            rejected = client.post("/system/user/import/data", files={"file": ("invalid.xlsx", filled_template(template.content, name + "x", "不存在的性别"), MIME)})
            assert rejected.status_code == 400
            message = rejected.json()["msg"]
            assert "第2行" in message and "性别" in message and "不存在的性别" in message
            assert data(client.get("/system/user/list", params={"username": name + "x"}))["total"] == 0
            print("PASS: invalid label rejected with row/field/value and no user persisted")

            for resource in ("role", "position", "ticket"):
                result = client.post(f"/system/{resource}/export", json={})
                if resource == "ticket" and result.status_code == 400:
                    assert "没有数据" in result.json()["msg"]
                    print("PASS: empty ticket export handled (no live ticket rows)")
                    continue
                result.raise_for_status()
                book = load_workbook(io.BytesIO(result.content))
                rows = list(book.active.values)
                if resource in ("role", "position"):
                    for values in rows[1:]:
                        row = dict(zip(rows[0], values, strict=True))
                        assert row["状态"] in ("启用", "停用")
                book.close()
                print(f"PASS: {resource} export, {len(rows) - 1} rows")
        finally:
            items = data(client.get("/system/user/list", params={"username": name}))["items"]
            ids = [item["id"] for item in items if item["username"] in (name, name + "x")]
            if ids:
                data(client.request("DELETE", "/system/user/delete", json=ids))
                # Deletion uses the existing request-scoped transaction. Verify cleanup boundedly.
                for _ in range(20):
                    if data(client.get("/system/user/list", params={"username": name}))["total"] == 0:
                        break
                else:
                    raise AssertionError("Temporary user cleanup not visible")
            print("PASS: temporary test users removed from active records")


if __name__ == "__main__":
    main()
