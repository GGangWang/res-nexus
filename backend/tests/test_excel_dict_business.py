"""Exercise actual database/enum dictionaries through business Excel entrypoints."""

import copy
import io
from unittest.mock import AsyncMock

import pytest
from fastapi import UploadFile
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.base_schema import AuthSchema
from app.core.exceptions import CustomException
from app.modules.system.dict.model import DictInfoModel
from app.modules.system.dict.schema import DictNodeWrite
from app.modules.system.dict.service import DictService
from app.modules.system.position.service import PositionService
from app.modules.system.role.service import RoleService
from app.modules.system.ticket.service import TicketService
from app.modules.system.user.service import UserService
from app.plugin.module_demo.test_demo_base.service import TestDemoBaseService as BaseDemoService
from app.plugin.module_example.demo.service import DemoService
from app.utils.excel_dict import ExcelDictConfig, ExcelDictSnapshot
from app.utils.excel_util import ExcelUtil


@pytest.fixture
async def dictionary_db():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(DictInfoModel.__table__.create)
    try:
        async with async_sessionmaker(engine, expire_on_commit=False)() as db:
            auth = AuthSchema()
            auth.user.is_superuser = True
            dictionary = DictService(auth, db)
            root = await dictionary.save(AsyncMock(), DictNodeWrite(dict_name="性别", dict_code="sys_user_sex"))
            for code, label in (("0", "男"), ("1", "女"), ("2", "未知")):
                await dictionary.save(AsyncMock(), DictNodeWrite(parent_id=root["id"], dict_code=code, dict_name=label))
            yield auth, db, dictionary
    finally:
        await engine.dispose()


async def test_database_and_enum_sources(dictionary_db):
    _, _, service = dictionary_db
    converter = await ExcelDictSnapshot.load(service, {"gender": ExcelDictConfig("sys_user_sex"), "status": ExcelDictConfig("sys_enable_status", code_type=int)}, importing=True)
    assert converter.import_row({"gender": "男", "status": "启用"}, 2) == ({"gender": "0", "status": 1}, [])
    await service.db.flush()


async def test_user_template_filled_and_imported(dictionary_db):
    auth, db, _ = dictionary_db
    service = UserService(auth, db)
    output = await service.get_import_template()
    wb = load_workbook(io.BytesIO(output))
    assert {wb["字典选项"].cell(index, 6).value for index in range(1, 4)} == {"男", "女", "未知"}
    assert wb["字典选项"]["G1"].value == "启用"
    wb.active.append([1, "excel_user", "测试用户", None, None, "女", "启用"])
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    service._process_import_row = AsyncMock(return_value=(1, None))
    result = await service.batch_import(UploadFile(io.BytesIO(stream.getvalue()), filename="users.xlsx"))
    assert "成功导入 1" in result
    row = service._process_import_row.await_args.args[1]
    assert row["gender"] == "1" and row["status"] == 1
    assert not isinstance(row["status"], str)

    invalid = await ExcelUtil.aexport_list2excel([{"gender": "错误性别", "status": "启用"}], {"gender": "性别", "status": "状态"})
    service._process_import_row.reset_mock()
    with pytest.raises(CustomException) as error:
        await service.batch_import(UploadFile(io.BytesIO(invalid), filename="invalid.xlsx"))
    assert "第2行" in str(error.value) and "性别" in str(error.value)
    service._process_import_row.assert_not_awaited()


@pytest.mark.parametrize(
    "service_type,data,expected",
    [
        (UserService, {"gender": "0", "status": 1, "is_superuser": True}, {"性别": "男", "状态": "启用", "是否超级管理员": "是"}),
        (RoleService, {"data_scope": 2, "status": 0}, {"数据权限": "本部门及以下数据权限", "状态": "停用"}),
        (PositionService, {"status": 1}, {"状态": "启用"}),
        (TicketService, {"status": 2, "ticket_type": "bug"}, {"工单状态": "已完成", "工单类型": "🐛 缺陷"}),
    ],
)
async def test_business_exports(dictionary_db, service_type, data, expected):
    auth, db, _ = dictionary_db
    original = copy.deepcopy(data)
    result = await service_type(auth, db).export_list([data])
    row = ExcelUtil.read_excel_to_dicts(result)[0]
    assert all(row[key] == value for key, value in expected.items())
    assert data == original


@pytest.mark.parametrize(
    "service_type,data,expected",
    [
        (DemoService, {"name": "示例", "status": 1}, {"状态": "启用"}),
        (BaseDemoService, {"user_name": "示例", "user_sex": 0}, {"性别": "男"}),
    ],
)
async def test_plugin_exports(dictionary_db, service_type, data, expected):
    auth, db, _ = dictionary_db
    original = copy.deepcopy(data)
    output = await service_type(auth, db).batch_export([data])
    row = ExcelUtil.read_excel_to_dicts(output)[0]
    assert all(row[key] == value for key, value in expected.items())
    assert data == original


@pytest.mark.parametrize(
    "service_type,values,crud_path,field,code",
    [
        (DemoService, ["示例", "启用", "备注"], "app.plugin.module_example.demo.service.DemoCRUD", "status", 1),
        (BaseDemoService, ["示例", "男"], "app.plugin.module_demo.test_demo_base.service.TestDemoBaseCRUD", "user_sex", 0),
    ],
)
async def test_plugin_template_to_import(dictionary_db, service_type, values, crud_path, field, code):
    from unittest.mock import patch

    auth, db, _ = dictionary_db
    service = service_type(auth, db)
    output = await service.import_template_download()
    wb = load_workbook(io.BytesIO(output))
    wb.active.append(values)
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    crud = AsyncMock()
    crud.get.return_value = None
    with patch(crud_path, return_value=crud):
        result = await service.batch_import(UploadFile(io.BytesIO(stream.getvalue()), filename="demo.xlsx"))
    assert "成功导入 1" in result
    schema = crud.create.await_args.kwargs["data"]
    assert getattr(schema, field) == code
