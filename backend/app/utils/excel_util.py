import asyncio
import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from app.common.datetime_utils import format_datetime
from app.utils.excel_dict import ExcelDictSnapshot, excel_raw


class ExcelUtil:
    """Excel 模板生成与列表导出（openpyxl）。"""

    @staticmethod
    def read_excel_to_dicts(contents: bytes, *, dictionary: ExcelDictSnapshot | None = None, header_mapping: dict | None = None) -> list[dict[str, Any]]:
        """读取 Excel 文件字节，返回字典列表（首行为列名）。"""
        wb = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
        try:
            ws = wb.active
            if not ws:
                raise ValueError("工作簿没有活动工作表")
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            result: list[dict[str, Any]] = []
            errors = []
            for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                if all(cell is None for cell in row):
                    continue
                row_dict: dict[str, Any] = {}
                for i, val in enumerate(row):
                    if i < len(headers) and headers[i] is not None:
                        row_dict[str(headers[i])] = val
                if row_dict:
                    if dictionary:
                        mapping = header_mapping or {h: h for h in headers if h is not None}
                        row_dict = {mapping.get(key, key): value for key, value in row_dict.items()}
                        row_dict, row_errors = dictionary.import_row(row_dict, row_number, {v: k for k, v in mapping.items()})
                        errors.extend(row_errors)
                    result.append(row_dict)
            if errors:
                raise ValueError("\n".join(errors))
            return result
        finally:
            wb.close()

    @staticmethod
    async def aread_excel_to_dicts(contents: bytes, **kwargs) -> list[dict[str, Any]]:
        """read_excel_to_dicts 的异步版本：解析在线程池中执行。

        openpyxl 解析上传文件（含 zip 解压与逐单元格构建）是纯同步 CPU +
        内存操作，大文件会在事件循环里独占数秒，async 调用点必须使用本方法。
        """
        return await asyncio.to_thread(ExcelUtil.read_excel_to_dicts, contents, **kwargs)

    @classmethod
    def __mapping_list(cls, list_data: list[dict[str, Any]], mapping_dict: dict) -> list[dict[str, Any]]:
        """将列表数据中的字段名映射为对应的中文字段名。

        参数:
        - list_data: 数据列表。
        - mapping_dict: 字段名映射字典 {英文key: 中文表头}。

        返回:
        - list[dict]: 映射后的数据列表 [{中文表头: value}]。
        """
        return [{str(mapping_dict.get(key)): item.get(key) for key in mapping_dict} for item in list_data]

    @classmethod
    def get_excel_template(
        cls,
        header_list: list[str],
        selector_header_list: list[str],
        option_list: list[dict[str, list[str]]],
        *,
        dictionary: ExcelDictSnapshot | None = None,
        mapping_dict: dict | None = None,
    ) -> bytes:
        """生成 Excel 模板文件。

        参数:
        - header_list: 表头列表。
        - selector_header_list: 需要设置下拉选择的表头列表。
        - option_list: 下拉选项配置列表。

        返回:
        - bytes: Excel 文件的二进制数据。
        """
        wb = Workbook()
        ws = wb.active
        if not ws:
            raise ValueError("不存在活动工作表")

        header_fill = PatternFill(start_color="ababab", end_color="ababab", fill_type="solid")

        for col_num, header in enumerate(header_list, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header  # pyright: ignore[reportAttributeAccessIssue]
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col_num)].width = 12

        for selector_header in selector_header_list:
            col_idx = header_list.index(selector_header) + 1
            header_options = next(
                (opt.get(selector_header) for opt in option_list if selector_header in opt),
                [],
            )
            if header_options:
                dv = DataValidation(type="list", formula1=f'"{",".join(header_options)}"')
                dv.add(f"{get_column_letter(col_idx)}2:{get_column_letter(col_idx)}1048576")
                ws.add_data_validation(dv)

        if dictionary:
            options_sheet = wb.create_sheet("字典选项")
            options_sheet.sheet_state = "hidden"
            for field, config in dictionary.configs.items():
                header = (mapping_dict or {}).get(field, field)
                if header not in header_list:
                    continue
                col_idx = header_list.index(header) + 1
                cell = ws.cell(1, col_idx)
                if config.mode == "single":
                    options = dictionary.options(field)
                    for index, label in enumerate(options, 1):
                        options_sheet.cell(index, col_idx, label)
                    if options:
                        letter = get_column_letter(col_idx)
                        name = f"excel_dict_{col_idx}"
                        wb.defined_names.add(DefinedName(name, attr_text=f"'字典选项'!${letter}$1:${letter}${len(options)}"))
                        dv = DataValidation(type="list", formula1=name)
                        dv.add(f"{letter}2:{letter}1048576")
                        ws.add_data_validation(dv)
                else:
                    example = dictionary.examples(field)
                    cell.comment = Comment(
                        f"支持名称 JSON 数组或分隔文本。多项分隔符: {config.separator!r}；层级分隔符: {config.path_separator!r}。示例: {example}",
                        "ResNexus",
                    )

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    def export_list2excel(cls, list_data: list[dict[str, Any]], mapping_dict: dict, *, dictionary: ExcelDictSnapshot | None = None) -> bytes:
        """将列表数据导出为 Excel 文件。

        参数:
        - list_data: 要导出的数据列表。
        - mapping_dict: 字段名映射字典 {英文key: 中文表头}。

        返回:
        - bytes: Excel 文件的二进制数据。

        限制:
        - 最多导出 100000 条记录，超出部分截断。
        """
        max_rows = 100000
        if len(list_data) > max_rows:
            list_data = list_data[:max_rows]

        if dictionary:
            list_data = [{key: dictionary.convert(key, value) if key in dictionary.configs else value for key, value in row.items()} for row in list_data]
        mapping_data = cls.__mapping_list(list_data, mapping_dict)
        headers = list(mapping_dict.values())

        wb = Workbook()
        ws = wb.active
        if not ws:
            raise ValueError("不存在活动工作表")

        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        for row_num, row_data in enumerate(mapping_data, 2):
            for col_num, header in enumerate(headers, 1):
                value = row_data.get(header)
                ws.cell(row=row_num, column=col_num, value=format_datetime(value) if isinstance(value, datetime) else excel_raw(value))

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    async def aexport_list2excel(cls, list_data: list[dict[str, Any]], mapping_dict: dict, **kwargs) -> bytes:
        """export_list2excel 的异步版本：构建与序列化在线程池中执行。

        导出最多 10 万行，逐单元格写入 + ZIP 压缩为纯同步 CPU 操作，
        千行以上即会明显卡住事件循环，async 调用点必须使用本方法。
        """
        return await asyncio.to_thread(cls.export_list2excel, list_data, mapping_dict, **kwargs)
